"""
Catálogo de las elecciones LOCALES de la Ciudad de México (IECM), como SQL que
se corre directo contra `electorama`, antes de load-results.sql.

Escribe, para las elecciones de `normalize.IECM`:
  · el partido local ELIGE (2021) y su bloque;
  · la fuente (archivo, url y sha256) y la elección con state_id = 9;
  · la boleta (ballot_option) con TODAS las columnas de votos del archivo;
  · las candidaturas por territorio (alcaldía en MAYOR, distrito local en
    LOCAL_DEPUTY, la ciudad en GOVERNOR).

Candidaturas: la misma regla que iepct_catalog.py. Una opción está en la
boleta de un territorio si su celda no viene vacía (2021 deja vacía la que no
estaba) y suma votos ahí (2024 escribe 0 en vez de vacío). Una combinación
presente ("PVEM_PT_MORENA") une a sus partidos en una sola candidatura; los
candidatos sin partido son IND1, IND2… en el orden de sus columnas.

Re-correrlo reemplaza esas elecciones completas: hay que volver a cargar los
resultados después.

Lee los st_/vt_ que dejó `normalize.py … --source iecm`. Uso (desde la carpeta
de las fuentes):
  python iecm_catalog.py <carpeta de normalizados> <salida.sql>
"""

import csv
import hashlib
import re
import sys
from collections import defaultdict
from pathlib import Path

from normalize import IECM, IECM_STATE

URL = {
    2024: 'https://estadisticaresultadospelo2024.iecm.mx/archivos/download{}.php',
    2021: 'https://www.iecm.mx/www/estadisticaresultadospelo2021/archivos/download{}.php',
}
DOWNLOAD = {'GOVERNOR': 'jg', 'LOCAL_DEPUTY': 'dmr', 'MAYOR': 'alc'}
SCOPE = {'GOVERNOR': 'STATE', 'LOCAL_DEPUTY': 'LOCAL_DISTRICT', 'MAYOR': 'MUNICIPALITY'}
DATE = {2021: '2021-06-06', 2024: '2024-06-02'}
DOWNLOADED = '2026-10-30'

PARTY = {'PAN', 'PRI', 'PRD', 'PVEM', 'PT', 'MC', 'MORENA', 'ELIGE', 'RSP', 'FXM'}
PES = {2021: 'PES2021'}
ORDER = ['PAN', 'PRI', 'PRD', 'PVEM', 'PT', 'MC', 'MORENA', 'ELIGE', 'PES2021', 'RSP', 'FXM']


def q(text) -> str:
    return "'" + str(text).replace("'", "''") + "'"


def arr(codes) -> str:
    return "'{" + ','.join(codes) + "}'"


def parties(label: str, year: int):
    """'PVEM_PT_MORENA' → ['PVEM', 'PT', 'MORENA']; None si no es de partidos (sin partido)."""
    codes = []
    for token in label.upper().split('_'):
        if token == 'PES':
            codes.append(PES[year])
        elif token in PARTY:
            codes.append(token)
        else:
            return None
    return sorted(codes, key=ORDER.index)


def independents(year: int, office: str) -> dict:
    """Sigla → nombre, de la hoja de abreviaturas ('JLSB: Candidato sin partido - José …')."""
    names = {}
    for row in iecm_rows_last_sheet(Path.cwd() / IECM[(year, office)]):
        cells = [str(c).strip() for c in row if c not in (None, '')]
        for i, cell in enumerate(cells):
            text = cell if ':' in cell else (f'{cell}: {cells[i + 1]}' if i + 1 < len(cells) else '')
            m = re.match(r'^([A-Z]+):\s*Candidat[oa] sin partido\s*-\s*([^(.]+)', text, re.I)
            if m:
                names[m.group(1).upper()] = ' '.join(m.group(2).split())
    return names


def iecm_rows_last_sheet(path: Path):
    if path.suffix == '.xls':
        import xlrd

        ws = xlrd.open_workbook(path).sheets()[-1]
        return [ws.row_values(i) for i in range(ws.nrows)]
    from openpyxl import load_workbook

    ws = load_workbook(path, read_only=True, data_only=True).worksheets[-1]
    return [list(r) for r in ws.iter_rows(values_only=True)]


def catalog(year: int, office: str, normalized: Path):
    territory = {}
    with open(normalized / f'st_{year}_{office}.txt', encoding='utf-8') as f:
        for r in csv.reader(f, delimiter='|'):
            key, district, mun = r[0], r[2], r[3]
            if office == 'MAYOR':
                territory[key] = mun if mun not in ('', '0') else ''
            elif office == 'LOCAL_DEPUTY':
                territory[key] = district if district != '0' else ''
            else:
                territory[key] = None

    labels, votes = [], defaultdict(int)
    with open(normalized / f'vt_{year}_{office}.txt', encoding='utf-8') as f:
        for key, label, n in csv.reader(f, delimiter='|'):
            if label not in labels:
                labels.append(label)
            t = territory[key]
            if t != '':  # extranjero y prisión no son de ningún distrito ni alcaldía
                votes[(t, label)] += int(n)

    names = independents(year, office)
    options = []
    for label in labels:
        codes = parties(label, year)
        if codes is None:
            code = f'IND{sum(1 for o in options if o[1] == "INDEPENDENT") + 1}'
            # 2024 nombra la columna CAND_IND1 y la abreviatura CSP: si sólo hay uno, es ése
            name = names.get(label.upper()) or (next(iter(names.values())) if len(names) == 1 else label)
            options.append((label, 'INDEPENDENT', [code], f'Sin partido ({name})'))
        else:
            options.append((label, 'PARTY' if len(codes) == 1 else 'COMBINATION', codes, '-'.join(codes)))
    option = {o[0]: o for o in options}

    present = defaultdict(set)
    for (t, label), n in votes.items():
        if n > 0:
            present[t].add(label)

    candidacies = []
    for t in sorted(present, key=lambda t: int(t) if t else 0):
        parent = {}

        def find(x):
            while parent[x] != x:
                parent[x] = parent[parent[x]]
                x = parent[x]
            return x

        for label in present[t]:
            kind, codes = option[label][1], option[label][2]
            if kind != 'INDEPENDENT':
                for c in codes:
                    parent.setdefault(c, c)
                for c in codes[1:]:
                    parent[find(c)] = find(codes[0])
        groups = defaultdict(list)
        for c in parent:
            groups[find(c)].append(c)
        for members in groups.values():
            members = sorted(members, key=ORDER.index)
            candidacies.append((t, '_'.join(members), '-'.join(members), members))
        for label in sorted(present[t]):
            if option[label][1] == 'INDEPENDENT':
                candidacies.append((t, option[label][2][0], option[label][3], option[label][2]))
    return options, candidacies


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    normalized, out = Path(sys.argv[1]), Path(sys.argv[2])
    sql = ['-- generado por tools/normalize/iecm_catalog.py: no editar a mano',
           '\\set ON_ERROR_STOP on', 'BEGIN;', '',
           "INSERT INTO electoral_history.party (code, name, color) VALUES "
           "('ELIGE', 'Partido Equidad, Libertad y Género', '#8E44AD') ON CONFLICT DO NOTHING;",
           "INSERT INTO electoral_history.bloc_party (year, party_code, bloc_code) VALUES "
           "(2021, 'ELIGE', 'OTHERS') ON CONFLICT DO NOTHING;", '']
    total = 0
    for (year, office), file_name in IECM.items():
        path = Path.cwd() / file_name
        sql.append(f'-- {office} {year}')
        sql.append(
            'INSERT INTO electoral_history.source (publisher, url, file_name, sha256, cut_label, downloaded_at) VALUES '
            f"('IECM', {q(URL[year].format(DOWNLOAD[office]))}, {q(file_name)}, {q(sha256(path))}, "
            f"{q('Estadística de resultados por casilla electoral del IECM')}, {q(DOWNLOADED)}) "
            'ON CONFLICT (file_name) DO UPDATE SET sha256 = EXCLUDED.sha256;')
        sql.append(
            'INSERT INTO electoral_history.election (year, office, scope, election_date, source_id, state_id) '
            f"SELECT {year}, {q(office)}, {q(SCOPE[office])}, {q(DATE[year])}, id, {IECM_STATE} "
            f'FROM electoral_history.source WHERE file_name = {q(file_name)} ON CONFLICT DO NOTHING;')
        election = (f'(SELECT id FROM electoral_history.election WHERE year = {year} '
                    f'AND office = {q(office)} AND state_id = {IECM_STATE})')
        sql.append(f'DELETE FROM electoral_history.station_vote WHERE station_id IN '
                   f'(SELECT id FROM electoral_history.polling_station WHERE election_id = {election});')
        for table in ('polling_station', 'official_total', 'candidacy', 'ballot_option'):
            sql.append(f'DELETE FROM electoral_history.{table} WHERE election_id = {election};')

        options, candidacies = catalog(year, office, normalized)
        sql.append('INSERT INTO electoral_history.ballot_option (election_id, raw_label, kind, party_codes, label) VALUES')
        sql.append(',\n'.join(f'  ({election}, {q(l)}, {q(k)}, {arr(c)}, {q(t)})' for l, k, c, t in options) + ';')
        sql.append('INSERT INTO electoral_history.candidacy (election_id, state_id, territory, code, name, party_codes) VALUES')
        sql.append(',\n'.join(
            f"  ({election}, {IECM_STATE}, {'NULL' if t is None else q(t)}, {q(code)}, {q(name)}, {arr(c)})"
            for t, code, name, c in candidacies) + ';')
        sql.append('')
        total += len(candidacies)
        print(f'{office} {year}: {len(options)} columnas de boleta, {len(candidacies)} candidaturas')
    sql.append('COMMIT;')
    out.write_text('\n'.join(sql) + '\n', encoding='utf-8', newline='\n')
    print(f'{out.name}: {total} candidaturas')


if __name__ == '__main__':
    main()
